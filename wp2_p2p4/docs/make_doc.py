"""Builds wp2_p2p4/docs/WP2_P2P4_Early_Warning.docx (python-docx)."""
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = sys.argv[1]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.0)
TEXT_W = 21.0 - 4.4

st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(5)
for name, size, color in (("Heading 1", 15, "1F3864"), ("Heading 2", 12.5, "2E5597"), ("Heading 3", 11, "2E5597")):
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
    h.font.color.rgb = RGBColor.from_string(color)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(12 if name == "Heading 1" else 9), Pt(4)


def P(text="", bold=False, italic=False, style=None, size=None):
    p = doc.add_paragraph(style=style)
    _runs(p, text, bold, italic, size)
    return p


def _runs(p, text, bold=False, italic=False, size=None):
    """**bold** and `code` inline markup."""
    import re
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        else:
            r = p.add_run(part); r.bold = bold
        r.italic = italic or r.italic
        if size:
            r.font.size = Pt(size)


def B(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        _runs(p, it)
        p.paragraph_format.space_after = Pt(2)


def N(items):
    B(items, "List Number")


def H(text, lvl=1):
    doc.add_heading(text, lvl)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def T(rows, widths=None, header=True, size=9):
    n = len(rows[0])
    widths = widths or [TEXT_W / n] * n
    t = doc.add_table(rows=len(rows), cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _runs(p, str(val), bold=(header and i == 0), size=size)
            if header and i == 0:
                shade(c, "D9E2F3")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ------------------------------------------------------------------ title
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = t.add_run("WP2 — Early warning for P2 / P4 output failures"); r.bold = True; r.font.size = Pt(20)
r.font.color.rgb = RGBColor.from_string("1F3864")
P("Sensata CFT project · Work Package 2 · analysis, design decisions and build log (steps a–e)", italic=True)
P("Author: Irfan Ramizu · Status: in progress, 22 October 2026 · Code: `wp2_p2p4/`", italic=True, size=9.5)

H("Summary")
P("WP2 builds a detector that warns when a CFT tester fixture (or a whole machine) is drifting towards "
  "its P2 or P4 output limit, ideally before the first reject. It reuses the design of the “Trace Sentinel” "
  "detector (streaming per-group features, machine-wide move removal, frozen centre, CUSUM, alarm tracker, "
  "false-alarm-budget tuning) on the WP1 dataset, which has real test times and CFT machine × fixture "
  "but no link to calibration.")
B(["**Data and limits.** The per-device limits file is not on disk, but WP1 used it; target and tolerance are "
   "recovered exactly from WP1’s deviation columns (83 P2 devices, 44 P4 devices, zero spread).",
   "**Reject rule resolved.** `|error| / tolerance > 1` reproduces the tester’s own P2/P4 output verdict with "
   "precision 0.992 / recall 0.883 (P2) and 0.996 / 0.998 (P4). The P2 misses are two devices whose limit "
   "changed during the period. The tester’s status is therefore used as the reject label.",
   "**Two corrections to the reference design** were needed on real data: units are not independent "
   "(CUSUM saturated on healthy data until a noise-inflation factor was added), and the “trend towards the "
   "limit” feature had the wrong sign for drifts on the negative side.",
   "**Blind spot confirmed.** Subtracting the machine-wide move cancels a drift that hits every fixture of a "
   "machine at once (e.g. an upstream calibration offset). A machine-level detector is planned (step g).",
   "**Steps a–e are built and checked on a sample** (3 machines, 1–14 June). On that sample, with untuned "
   "thresholds, CUSUM and the trained trees caught 9 of 10 injected drifts, all before the first reject, where the "
   "count rule caught 2. Evaluation on full data, whole-machine drift and export (f–h) follow."])

# ------------------------------------------------------------------ 1
H("1. Goal, data and rules")
P("P2 and P4 are **output tests**: the sensor’s output at a test pressure is compared with the expected "
  "output. They are not leak tests. The aim is early warning: notice that a fixture’s passing units are "
  "creeping towards the limit before they cross it.")
H("1.1 Data (WP1 output, read-only)", 2)
T([["Item", "Value"],
   ["Location", "`wp1_data/wp1_data/explain/{train,test}/part-0000.parquet`"],
   ["Rows", "train 4,255,223 · test 897,021 · 94 columns"],
   ["Train period", "1 May – 23 Jul 2026 (test starts at the 23 Jul cut, no shared serials)"],
   ["First-pass tests in train", "3,964,099 (93 % of rows have attempt_no = 1)"],
   ["Machines / devices", "20 CFT machines (equipmentid) · 90 devices · tester fixtures 1–10"],
   ["P2 / P4 output reject base rate (train)", "0.21 % of P2-tested units · 0.06 % of P4-tested units"]],
  [5.0, TEXT_W - 5.0])
H("1.2 Rules followed throughout", 2)
B(["Never re-split train/test. Validation is carved from the end of train, by time, with serials disjoint.",
   "Never report accuracy; always quote the base rate. Use precision/recall, PR-AUC, lead time and false "
   "alarms per machine per day.",
   "Null is not zero: 65535 / 9999 / 999 / −999 and output = 0 are placeholders, not readings.",
   "First attempts only (attempt_no = 1) for baselines.",
   "No real serial numbers in reports or outputs (a salted hash `unit_key` is used instead)."])

# ------------------------------------------------------------------ 2
H("2. Earlier WP2 work (inventory)")
T([["File", "What it does", "Inputs → outputs"],
   ["tests_codes/near_threshold_v2.py", "Flags passing units in the top 1 % of their machine × device × fixture "
    "distribution on the LEAK channels P1/P3/P5, plus an Isolation Forest on 18 *_dev columns. Static "
    "baseline fitted once on train. Validated on the leak-fail rate of the next 50 tests, stratified by "
    "fx_recent_fail.", "explain train/test → scored test units CSV"],
   ["tests_codes/near_threshold_v3.py", "v2 plus an EWMA/CUSUM drift monitor per machine + fixture + leak channel, "
    "an alert-budget table and a lead-time check against chance.", "same → scored CSV, drift alerts CSV, model bundle"],
   ["tests_codes/make_synthetic_split.py", "Synthetic leak data with one drifting fixture, for testing.", "→ synth CSVs"],
   ["notepad.txt", "Console logs of the v2/v3 runs (status tables only).", "—"]],
  [4.3, 8.7, TEXT_W - 13.0])
P("Results reported from those runs: about 6.9 % of passing units flagged (expected 1–3 %), lift of about "
  "1.0–1.3 over the base rate (no real signal), and machine 9990675 flagged on every fixture. Note that this "
  "work targeted leak failures; P2/P4 output early warning is new.")

# ------------------------------------------------------------------ 3
H("3. The reference design: Trace Sentinel")
P("Built on the September 2026 extract (no timestamps, calibration machine × fixture group n/n+4/n+8) and "
  "evaluated on simulated drifts.")
B(["**Signal:** signed error in % of tolerance, `v = 100 × error / OutputTol`; “not measured” when output "
   "and error are both 0; first pass only. One stream per (calibration machine, fixture group, signal).",
   "**Warm-up:** the first 120 values set the group’s centre (median); noise σ from their MAD × 1.4826, then "
   "updated online and shared across the machine’s groups.",
   "**Machine-wide move removal:** the median move of the other warm groups away from their own centres is "
   "subtracted, so machine wander and lot changes cancel. The median stops one drifting neighbour from "
   "dragging the reference.",
   "**Frozen centre:** the centre learns slowly (span 1,500) only while the group looks normal "
   "(|d_m| < 0.35σ and CUSUM < 4) and freezes while it looks displaced; a shift lasting over 600 units "
   "becomes the new centre.",
   "**CUSUM:** two-sided, k = 0.5, z clipped to ±4, capped at 40; the last time it sat at zero is the "
   "estimated drift onset.",
   "**Alarm tracker:** ok → warning → alarm after 3 consecutive scores over threshold; an alarm is latched; "
   "clears after 60 scores below 0.8 × the warning threshold.",
   "**Labels:** injected drift episodes (unit moved ≥ 10 % tol by the drift); on real data, the run-up to a "
   "reject cluster (≥ 3 rejects in the next 200 units).",
   "**Evaluation:** time split 55 / 20 / 25 % (train / tune / test). For each false-alarm budget "
   "(0.1, 0.25, 0.5, 1, 2 per machine per day) the lowest thresholds within budget are found on tune; test "
   "is scored once. Baselines: count rule (3 rejects in 200 units) and a z-rule."])
P("The 15 features:", bold=True)
T([["Feature", "Meaning"],
   ["d_s, d_m, d_l", "Departure of the 30 / 100 / 300-unit average from the centre, minus the machine move, in σ"],
   ["abs_d_s, abs_d_m", "Absolute short and medium departure"],
   ["trend", "Short minus long average: drift per 100 units, towards the limit"],
   ["toward", "+1 if the departure points towards the nearer limit, −1 otherwise"],
   ["cusum", "max of the two CUSUM sides"],
   ["headroom", "(100 − |level|) / 100: room left to the limit"],
   ["eta_log", "log10 of projected units to the limit (headroom / slope)"],
   ["near_g", "Rate of units with 80 ≤ |v| ≤ 100"],
   ["near_diff, rej_diff", "Near-limit and reject rate minus the other groups’ average"],
   ["rej_w200", "Rejects in the last 200 units (the count rule’s input)"],
   ["warm", "Warm-up progress, 0–1"]],
  [3.6, TEXT_W - 3.6])

# ------------------------------------------------------------------ 4
H("4. Gap analysis")
T([["Aspect", "Earlier WP2 (v2/v3)", "Trace Sentinel", "Consequence"],
   ["Target", "Leak P1/P3/P5, one-sided", "P2/P4 output, two-sided", "Not directly comparable"],
   ["Grouping", "machine × device × fixture, fallback device × fixture", "cal machine × fixture group",
    "WP2 groups are finer"],
   ["Normalisation", "Percentile → normal score per group", "Signed % of tolerance", "% of tolerance has a physical meaning"],
   ["Baseline", "Static, fitted once on train", "Online, frozen when displaced", "Static baselines flag every later level change"],
   ["Machine-wide move", "None", "Median of other groups subtracted", "Explains 9990675 flagged on every fixture"],
   ["Features", "Per-unit tail + Isolation Forest; EWMA/CUSUM", "15 streaming group features",
    "One unit’s tail predicts little → lift ~1"],
   ["Labels", "Leak-fail rate in next 50 tests", "Drift episodes / reject run-up", "Diluted by sporadic failures"],
   ["Validation", "Fixed cut-offs, no tune slice", "Tune slice, false-alarm budgets, test once",
    "Explains 6.9 % flagged instead of 1–3 %"],
   ["Outputs", "Unit flags, alert CSV", "Alert events with level, shift, ETA, onset", "Follow Trace Sentinel for Tab 1"]],
  [2.6, 4.3, 4.3, TEXT_W - 11.2], size=8.5)
H("4.1 Data differences", 2)
T([["Trace Sentinel (September extract)", "WP2 (WP1 data)"],
   ["Calibration machine × fixture group (n/n+4/n+8)", "CFT equipmentid × tester fixture 1–10"],
   ["No timestamps; file order as sequence", "Real testtime"],
   ["Explicit limits table, DoP2Test flag", "Limits recovered from WP1 columns; p2_tested flag"],
   ["Link to calibration", "No link to calibration"]],
  [TEXT_W / 2, TEXT_W / 2])
B(["**Ported directly:** % of tolerance, the not-measured rule, the streaming state, frozen centre, CUSUM, the "
   "15 features, alarm tracker, reject-cluster labels, budget tuning, count rule.",
   "**Adapted:** the peer set for the machine move (9 other tester fixtures on the same machine and device), "
   "“per day” in real time, the reject definition (step b), no recalibration resets.",
   "**Not portable:** back-trace to calibration and rework lists by calibration time.",
   "**Key implication:** if P2/P4 faults originate in calibration, units reach tester fixtures more or "
   "less at random, so a calibration drift appears on a CFT machine as a whole-machine (or whole-device) "
   "move — exactly what the per-fixture design subtracts. Per-fixture streams on this data detect tester "
   "fixture faults; a machine-level stream is needed for upstream faults (step g)."])

# ------------------------------------------------------------------ 5
H("5. Plan (approved)")
T([["Step", "Content", "Status"],
   ["a. Loader", "Signed P2/P4 error as % of tolerance, not-measured/placeholder rules, first pass only", "Done"],
   ["b. Reject-rule check", "Does |error| > tolerance match the tester’s P2/P4 output status?", "Done"],
   ["c. Features", "Port of features.py; key (machine, device, fixture, signal); machine move over the other fixtures", "Done"],
   ["d. Labels", "Reject-cluster run-up + drift injected into real healthy data at a known time", "Done"],
   ["e. Detectors", "Count rule (baseline), CUSUM, gradient-boosted trees, alarm tracker", "Done"],
   ["f. Evaluation", "Time split, thresholds tuned to 0.1–2 false alarms per machine per day, test scored once", "Next"],
   ["g. Whole-machine drift", "Inject drift on every fixture; add a machine-level stream", "Planned"],
   ["h. Export", "Scored CSV and alerts CSV for a Streamlit Tab 1", "Planned"]],
  [3.4, TEXT_W - 5.2, 1.8])
P("Evaluation split: fit on train up to about 2 July; tune on the last ~3 weeks of train (serials seen in "
  "the fit part removed); the WP1 test split is scored once at the end.")

# ------------------------------------------------------------------ 6
H("6. Build log")
H("6.1 Step a — loader (load.py)", 2)
P("One row per first-pass test with: `unit_key`, time, machine, device, fixture, failed, status family, and "
  "per signal `v_*` (% of tolerance), `why_*` (why v is set or empty), `fail_*` (tester output reject), "
  "`kind_*` (out_of_band / no_reading) and `band_ok_*`.")
B(["**Limits recovered.** WP1 computed `p2_rel_dev = (output − target) / target` and "
   "`p2_dev_tol = (output − target) / tol`. Solving back gives exactly one target and tolerance per device "
   "(e.g. 5MPP2-16 P2: 80.11 ± 0.5). Saved as `out/limits.csv`.",
   "**Status families.** P2/P4 output failures (Error Percent, V_Out High/Low, Error Invalid, High Rail) are "
   "separated from pressure-control failures (Stability, Fill Timeout, Mensor, Pressure Tolerance), which "
   "are pneumatic and not used as output rejects.",
   "**Aborted readings.** Units that failed an earlier step (mostly P1 leak, P1 stability, zero-point fault) "
   "still carry P2/P4 values, but these are readings without test pressure (≈ 660 counts vs ≈ 3,253 normal on "
   "62MPP03-04, up to −13,000 % of tolerance). On a failed unit, v is kept only if the unit failed that "
   "signal’s own output test or a later step; otherwise `why = aborted`."])
P("Sample (3 machines, 1–14 June, 127,410 first-pass tests):", bold=True)
T([["", "P2", "P4"],
   ["Measured (v set)", "114,388", "98,082"],
   ["Not tested / aborted / output = error = 0 / placeholder", "10,394 / 1,575 / 1,047 / 6", "26,349 / 1,568 / 1,400 / 11"],
   ["Healthy spread, 1st–99th percentile (% tol)", "−39 … +48", "−26 … +56 (centred near +23)"],
   ["Output rejects (status)", "26 (0.02 %)", "27 (0.02 %)"]],
  [7.0, 4.8, TEXT_W - 11.8])

H("6.2 Step b — reject rule (check_reject_rule.py, all of train)", 2)
P("Truth is the tester’s own P2/P4 output status. Rules are scored on units with a usable reading.")
T([["Signal / rule", "TP", "FP", "FN", "Precision", "Recall"],
   ["P2: |error| / tol > 1", "6,679", "52", "884", "0.992", "0.883"],
   ["P2: |output − target| / tol > 1", "6,663", "2,062,005", "900", "0.003", "0.881"],
   ["P4: |error| / tol > 1", "1,243", "5", "3", "0.996", "0.998"],
   ["P4: |output − target| / tol > 1", "1,088", "917,852", "158", "0.001", "0.873"]],
  [6.0, 1.8, 2.4, 1.6, 2.3, TEXT_W - 14.1])
B(["No passing unit is ever beyond ±100 % of tolerance; all false positives are units that failed for another reason.",
   "Measuring against the nominal target (what `p*_dev_tol` does) is not the tester’s rule: the error is "
   "relative to the expected output at the applied pressure.",
   "**The P2 misses come from two devices whose limit changed:** 31MPP3-2 (rejected at 64–119 % of the snapshot "
   "tolerance on the high side, 1,003 rejects in 2,525 units: a real production episode) and 85MPP36-01 "
   "(limit ≈ 80 % until the week of 8 June, then 100 %). All other devices separate exactly at 100 %.",
   "534 P2 and 117 P4 output rejects have no usable reading (output 0 or 65535): dead or saturated sensors "
   "rather than drift."])
P("Decision:", bold=True)
N(["Reject label = the tester’s status (not |v| > 100).",
   "v = 100 × error / tol remains the continuous signal; the two device periods are flagged `band_ok = False` "
   "and left out of headroom-based tuning and evaluation.",
   "Rejects are typed `out_of_band` (has a reading) or `no_reading`; reject-cluster labels and the count rule "
   "use out_of_band by default, both are reported.",
   "Open question for engineering: were the limits of 31MPP3-2 and 85MPP36-01 revised during May–July?"])

H("6.3 Step c — streaming features (features.py)", 2)
P("A port of Trace Sentinel’s features.py with the 15 features and the same maths, keyed by "
  "(machine, device, tester fixture, signal). Changes:")
B(["Machine-wide move = median move of the other warm fixtures of the same machine **and device**; σ shared "
   "per machine × device × signal.",
   "Rejects come from the tester status (out-of-band only).",
   "No recalibration records exist, so a fixture idle for 12 h (device changeover) restarts its averages "
   "and CUSUM at its centre.",
   "**Noise inflation (new).** On real data, consecutive units are not independent: the residual after "
   "centre and machine move has lag-1 autocorrelation ≈ 0.3 and 50-unit averages wander 0.4–4.7× more than "
   "independent units would (median 2.4×). With per-unit z, CUSUM sat above 10 for 34 % of healthy time. "
   "Each machine × device × signal now learns `infl` (observed / independent spread of a fixture’s 50-unit "
   "residual average, learnt only while the fixture looks normal, clamped 1–6) and CUSUM uses z / infl. "
   "Healthy CUSUM is now above 10 only 0.6 % of the time and never at the cap.",
   "**Trend sign (fix).** The reference multiplies trend by `toward`, which makes a drift moving down towards "
   "−100 look like it moves away. The port uses trend × sign(level), so positive always means “towards the "
   "nearer limit”."])
P("Checks on the sample (9990771 / 62MPP03-04 / P2, +40 % tol step from mid-window):", bold=True)
T([["Case", "Rows", "Shift (% tol)", "Departure d_m (σ)", "CUSUM 95th pct"],
   ["Step on one fixture", "that fixture, after", "+15", "1.54", "40"],
   ["", "other fixtures, after", "0", "0.00", "0.06"],
   ["Step on every fixture", "all fixtures, after", "≈ 0", "≈ 0", "≈ 24 (brief transient)"]],
  [3.6, 3.9, 2.6, 3.2, TEXT_W - 13.3])
B(["A one-fixture drift is seen on that fixture only. The average shift is +15 rather than +40 because a shift "
   "that lasts over 600 units is adopted as the new centre (by design).",
   "**A whole-machine drift is cancelled completely** by the machine-move removal: the blind spot step g addresses.",
   "Speed: about 180 µs per value, i.e. roughly 25 minutes for all of train and test; features will be cached."])

H("6.4 Step d — labels (labels.py)", 2)
B(["**Reject-cluster label (real data):** a row is positive when its stream has ≥ 3 out-of-band rejects among its "
   "next 200 units. On the sample only 2 of 188 streams have any positive (one fixture of 20190405 / 81MPP18-01): "
   "0.22 % of P2 rows and 0.28 % of P4 rows. Real reject clusters are rare, so they cannot be the only label.",
   "**Injected drift (known truth):** a step (30–70 % tol) or ramp (6–28 % tol per hour, capped at 170) is added "
   "to real healthy data from a known time, towards the side the stream already leans to. Episodes are drawn on "
   "warm streams with no real reject from 24 h before to 14 h after the start, at least 24 h apart per machine. "
   "Injected values beyond ±100 become rejects; after a P2 reject P4 is not measured (as on the tester). An "
   "episode ends after 5 drift rejects (“recalibrated”) or 14 h. Label = 1 while the drift has moved the unit "
   "≥ 10 % tol. Scope can be one fixture or every fixture of a machine × device (step g)."])
P("Sample result (seed 0): 13 one-fixture episodes, 8 reached a reject, 6 ended by the 5-reject rule, "
  "32 injected rejects; 1,165 positive rows of 188,443 (base rate 0.62 %). The first injected reject came "
  "21–123 units after the start.", italic=False)
T([["Row type (median)", "d_m (σ)", "|d_m|", "trend", "cusum", "headroom", "rej_w200"],
   ["Drift (label 1)", "−1.26", "2.22", "+1.55", "40", "0.71", "0"],
   ["Episode, moved < 10 %", "−0.01", "0.09", "+0.14", "0", "0.89", "0"],
   ["Healthy", "0.00", "0.06", "0.00", "0", "0.88", "0"]],
  [4.6, 2.0, 1.8, 1.8, 1.8, 2.1, TEXT_W - 14.1])
P("The features separate drifting rows clearly while rej_w200 (the count rule’s input) is still 0, which is "
  "the room an early-warning model has to beat the count rule. (Negative d_m reflects that 9 of 13 sample "
  "episodes drifted towards the lower limit.)")

H("6.5 Step e — detectors (detectors.py)", 2)
P("Port of Trace Sentinel’s detectors.py with unchanged maths; the alarm key is (machine, device, fixture, "
  "signal). The reference’s z-rule is not ported (it was only its replay demo’s stand-in).")
T([["Detector", "Score", "Default thresholds (warn / alarm)"],
   ["Count rule (baseline)", "rejects in the fixture’s last 200 units ÷ 3", "— / 1.0, no persistence"],
   ["CUSUM", "CUSUM × (moving towards the limit) × warm-up", "6 / 10, 3 consecutive"],
   ["Trained", "gradient-boosted trees (250 trees, 31 leaves, lr 0.06, positives weighted ≤ 20×)", "0.5 / 0.8, 3 consecutive"]],
  [3.6, 8.0, TEXT_W - 11.6])
P("Check on the sample: injected one-fixture drifts (19 episodes); trees trained on week 1 (1–7 June), all "
  "detectors scored on week 2 (8–14 June) at the default thresholds. Rows of real reject clusters that no "
  "injected episode explains are left out of training (they may be real drifts).")
T([["Week 2 (10 episodes, 3 machines × 7 days)", "Count rule", "CUSUM", "Trained"],
   ["Row-level PR-AUC (base rate 0.0145)", "0.015", "0.948", "0.980"],
   ["Episodes detected (alarm)", "20 %", "90 %", "90 %"],
   ["Alarm before the first reject (3 episodes with rejects)", "0 %", "100 %", "100 %"],
   ["Median delay, units of that fixture", "47", "11", "5"],
   ["False alarms / machine / day", "0.05", "0.48", "0.05"],
   ["False warnings / machine / day", "0.00", "0.76", "0.00"]],
  [7.4, 2.9, 2.9, TEXT_W - 13.2])
B(["The count rule acts after the damage (it needs 3 rejects), as Trace Sentinel found on simulated data.",
   "**These numbers are optimistic:** 10 test episodes on 3 machines, the same machines in training and "
   "scoring, and injected drifts of 30–70 % tol or 6–28 % tol per hour are large next to the noise "
   "(σ ≈ 3–12 % tol). Step f uses all machines, tunes thresholds on a separate slice to false-alarm budgets, "
   "adds smaller drifts to find the detection limit, and scores the test split once.",
   "The features most related to the label are the short departure (|d_s|), CUSUM, the medium departure "
   "and the trend towards the limit; the reject-count features carry almost nothing before the first reject."])

# ------------------------------------------------------------------ 7
H("7. Open questions and assumptions")
B(["Real DPSCFTLimits file, ideally with revision history (confirms the two band exceptions).",
   "Do P2/P4 faults originate in calibration (as Trace Sentinel assumed)? Then the machine-level stream matters most.",
   "Drift shapes and rates are Trace Sentinel’s simulator assumptions; no record of real drifts exists. The "
   "31MPP3-2 episode (40 % rejects for two weeks) is a candidate real episode to validate against.",
   "The 12 h changeover restart and the 5-reject recalibration rule are assumptions to confirm with the floor.",
   "Limits apply per device; P4 rejects are scarce (≈ 1,400 in train), so P4 evaluation relies on injected drift."])

# ------------------------------------------------------------------ 8
H("8. Files and how to run")
T([["File (wp2_p2p4/)", "Purpose", "Run"],
   ["load.py", "Loader, limits recovery, status families, aborted readings, band exceptions", "`python load.py`"],
   ["check_reject_rule.py", "Step b confusion tables", "`python check_reject_rule.py --machines \"\" --start \"\" --end \"\"`"],
   ["features.py", "Streaming feature engine (port of Trace Sentinel)", "library"],
   ["features_demo.py", "Step c checks (healthy data, one-fixture and whole-machine step)", "`python features_demo.py`"],
   ["labels.py", "Reject-cluster label, drift episode generator and injector", "library"],
   ["labels_demo.py", "Step d checks", "`python labels_demo.py`"],
   ["detectors.py", "Count rule, CUSUM, trained trees, alarm tracker", "library"],
   ["detectors_demo.py", "Step e check (train week 1, score week 2)", "`python detectors_demo.py`"],
   ["out/", "Generated files (limits.csv, logs, salt); git-ignored", "—"]],
  [3.8, 6.8, TEXT_W - 10.6], size=8.5)
P("Interpreter: `.venv/Scripts/python.exe` (populate with `uv sync --frozen`). Run from `wp2_p2p4/`.")

zoom = doc.settings.element.find(qn("w:zoom"))      # python-docx template omits the required percent
if zoom is not None:
    zoom.set(qn("w:percent"), "100")
doc.save(OUT)
print("saved", OUT)
